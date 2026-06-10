#!/usr/bin/env python3
"""
Automated test script for all required file conversions.
Tests every conversion from the requirements matrix against the running Flask app.
"""

import requests
import os
import sys
import json

BASE_URL = "http://localhost:5000"
TEST_DIR = os.path.dirname(os.path.abspath(__file__))
RESULTS = []

def create_test_files():
    """Create sample files for testing"""
    # TXT file
    with open(os.path.join(TEST_DIR, "test.txt"), "w") as f:
        f.write("Hello World - This is a test document.\nLine 2\nLine 3\n")
    
    # CSV file
    with open(os.path.join(TEST_DIR, "test.csv"), "w") as f:
        f.write("name,age,city\nJohn,25,NYC\nJane,30,LA\n")

    # Create a simple DOCX
    try:
        from docx import Document
        doc = Document()
        doc.add_paragraph("Test Document Content")
        doc.save(os.path.join(TEST_DIR, "test.docx"))
    except ImportError:
        print("⚠️  python-docx not available, skipping DOCX creation")

    # Create a simple PDF
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.pdfgen import canvas
        c = canvas.Canvas(os.path.join(TEST_DIR, "test.pdf"), pagesize=letter)
        c.drawString(100, 750, "Test PDF Document")
        c.save()
    except ImportError:
        print("⚠️  reportlab not available, skipping PDF creation")

    # Create a simple PNG image
    try:
        from PIL import Image
        img = Image.new('RGB', (100, 100), color='red')
        img.save(os.path.join(TEST_DIR, "test.png"))
        img.save(os.path.join(TEST_DIR, "test.jpg"), "JPEG")
        img.save(os.path.join(TEST_DIR, "test.jpeg"), "JPEG")
    except ImportError:
        print("⚠️  Pillow not available, skipping image creation")

    # Create a simple XLSX
    try:
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(["Name", "Age", "City"])
        ws.append(["John", 25, "NYC"])
        ws.append(["Jane", 30, "LA"])
        wb.save(os.path.join(TEST_DIR, "test.xlsx"))
    except ImportError:
        print("⚠️  openpyxl not available, skipping XLSX creation")

    # Create a simple PPTX
    try:
        from pptx import Presentation
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = "Test Presentation"
        slide.placeholders[1].text = "Test content"
        prs.save(os.path.join(TEST_DIR, "test.pptx"))
    except ImportError:
        print("⚠️  python-pptx not available, skipping PPTX creation")


def test_conversion(from_ext, to_ext, description=""):
    """Test a single conversion"""
    filepath = os.path.join(TEST_DIR, f"test.{from_ext}")
    
    if not os.path.exists(filepath):
        result = {"from": from_ext.upper(), "to": to_ext.upper(), "status": "⏭️ SKIP", "reason": f"No test.{from_ext} file", "desc": description}
        RESULTS.append(result)
        return result

    try:
        with open(filepath, 'rb') as f:
            files = {'file': (f"test.{from_ext}", f)}
            data = {'to_format': to_ext}
            response = requests.post(f"{BASE_URL}/api/convert", files=files, data=data, timeout=30)

        if response.status_code == 200:
            content_type = response.headers.get('content-type', '')
            size = len(response.content)
            if size > 0 and 'json' not in content_type:
                result = {"from": from_ext.upper(), "to": to_ext.upper(), "status": "✅ PASS", "reason": f"Got {size} bytes", "desc": description}
            else:
                try:
                    err = response.json()
                    result = {"from": from_ext.upper(), "to": to_ext.upper(), "status": "❌ FAIL", "reason": err.get('error', 'Empty response'), "desc": description}
                except:
                    result = {"from": from_ext.upper(), "to": to_ext.upper(), "status": "❌ FAIL", "reason": "Empty/invalid response", "desc": description}
        else:
            try:
                err = response.json()
                reason = err.get('error', f'HTTP {response.status_code}')
            except:
                reason = f"HTTP {response.status_code}: {response.text[:100]}"
            result = {"from": from_ext.upper(), "to": to_ext.upper(), "status": "❌ FAIL", "reason": reason, "desc": description}

    except requests.exceptions.ConnectionError:
        result = {"from": from_ext.upper(), "to": to_ext.upper(), "status": "🔌 OFFLINE", "reason": "Server not running", "desc": description}
    except Exception as e:
        result = {"from": from_ext.upper(), "to": to_ext.upper(), "status": "❌ ERROR", "reason": str(e)[:80], "desc": description}

    RESULTS.append(result)
    return result


def main():
    print("=" * 70)
    print("🧪 COMPREHENSIVE FILE CONVERSION TEST")
    print("=" * 70)
    
    # Check server is running
    try:
        r = requests.get(BASE_URL, timeout=5)
        print(f"✅ Server is running at {BASE_URL}")
    except:
        print(f"❌ Server is NOT running at {BASE_URL}")
        print("   Start with: python main.py")
        sys.exit(1)

    print("\n📁 Creating test files...")
    create_test_files()
    print("✅ Test files created\n")

    # ===== ALL REQUIRED CONVERSIONS =====
    
    print("=" * 70)
    print("📄 TEXT FILES (.TXT)")
    print("-" * 70)
    for to_fmt, desc in [("pdf", "TXT→PDF"), ("docx", "TXT→DOCX"), ("doc", "TXT→DOC"), ("html", "TXT→HTML")]:
        r = test_conversion("txt", to_fmt, desc)
        print(f"  {r['status']}  {r['from']} → {r['to']}  |  {r['reason']}")

    print("\n" + "=" * 70)
    print("📕 WORD DOCUMENTS (.DOCX / .DOC)")
    print("-" * 70)
    for from_fmt, to_fmt, desc in [
        ("docx", "pdf", "DOCX→PDF"), ("docx", "txt", "DOCX→TXT"), ("docx", "doc", "DOCX→DOC"),
        ("doc", "pdf", "DOC→PDF"), ("doc", "docx", "DOC→DOCX"), ("doc", "txt", "DOC→TXT"),
    ]:
        r = test_conversion(from_fmt, to_fmt, desc)
        print(f"  {r['status']}  {r['from']} → {r['to']}  |  {r['reason']}")

    print("\n" + "=" * 70)
    print("📕 PDF FILES (.PDF)")
    print("-" * 70)
    for to_fmt, desc in [("docx", "PDF→DOCX"), ("txt", "PDF→TXT"), ("png", "PDF→PNG"), ("jpg", "PDF→JPG"), ("doc", "PDF→DOC")]:
        r = test_conversion("pdf", to_fmt, desc)
        print(f"  {r['status']}  {r['from']} → {r['to']}  |  {r['reason']}")

    print("\n" + "=" * 70)
    print("🖼️  IMAGE FILES (.PNG / .JPG / .JPEG)")
    print("-" * 70)
    for from_fmt, to_fmt, desc in [
        ("png", "jpg", "PNG→JPG"), ("png", "pdf", "PNG→PDF"), ("png", "webp", "PNG→WEBP"),
        ("jpg", "png", "JPG→PNG"), ("jpg", "pdf", "JPG→PDF"), ("jpg", "webp", "JPG→WEBP"),
        ("jpeg", "png", "JPEG→PNG"), ("jpeg", "jpg", "JPEG→JPG"),
    ]:
        r = test_conversion(from_fmt, to_fmt, desc)
        print(f"  {r['status']}  {r['from']} → {r['to']}  |  {r['reason']}")

    print("\n" + "=" * 70)
    print("📊 SPREADSHEET FILES (.XLSX / .CSV)")
    print("-" * 70)
    for from_fmt, to_fmt, desc in [
        ("xlsx", "csv", "XLSX→CSV"), ("xlsx", "pdf", "XLSX→PDF"), ("xlsx", "txt", "XLSX→TXT"),
        ("csv", "xlsx", "CSV→XLSX"), ("csv", "pdf", "CSV→PDF"), ("csv", "txt", "CSV→TXT"), ("csv", "json", "CSV→JSON"),
    ]:
        r = test_conversion(from_fmt, to_fmt, desc)
        print(f"  {r['status']}  {r['from']} → {r['to']}  |  {r['reason']}")

    print("\n" + "=" * 70)
    print("🎬 PRESENTATION FILES (.PPTX)")
    print("-" * 70)
    for to_fmt, desc in [("pdf", "PPTX→PDF"), ("png", "PPTX→PNG"), ("jpg", "PPTX→JPG"), ("txt", "PPTX→TXT")]:
        r = test_conversion("pptx", to_fmt, desc)
        print(f"  {r['status']}  {r['from']} → {r['to']}  |  {r['reason']}")

    # ===== SUMMARY =====
    print("\n" + "=" * 70)
    print("📊 SUMMARY")
    print("=" * 70)
    passed = sum(1 for r in RESULTS if "PASS" in r['status'])
    failed = sum(1 for r in RESULTS if "FAIL" in r['status'] or "ERROR" in r['status'])
    skipped = sum(1 for r in RESULTS if "SKIP" in r['status'])
    total = len(RESULTS)
    
    print(f"  ✅ Passed:  {passed}/{total}")
    print(f"  ❌ Failed:  {failed}/{total}")
    print(f"  ⏭️  Skipped: {skipped}/{total}")
    print()
    
    if failed > 0:
        print("❌ FAILED CONVERSIONS:")
        for r in RESULTS:
            if "FAIL" in r['status'] or "ERROR" in r['status']:
                print(f"  {r['from']} → {r['to']}: {r['reason']}")


if __name__ == "__main__":
    main()
